#!/usr/bin/env python3
"""
embed_diagrams_in_presentation.py
Embeds the actual SRS diagrams, ERDs, architectural schemas, and GUI mockups
provided by the user directly into the PowerPoint presentation template.
Preserves the exact VU slide theme, master background, logo, and structure.
"""

import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def fit_image_box(img_path, box_left, box_top, box_w, box_h):
    """Calculates centered position and dimensions preserving aspect ratio."""
    im = Image.open(img_path)
    orig_w, orig_h = im.size
    aspect = orig_w / orig_h
    box_aspect = box_w / box_h

    if aspect > box_aspect:
        w = box_w
        h = box_w / aspect
        left = box_left
        top = box_top + (box_h - h) / 2
    else:
        h = box_h
        w = box_h * aspect
        top = box_top
        left = box_left + (box_w - w) / 2
    return left, top, w, h

def build_presentation_with_diagrams(input_pptx, output_pptx, updir):
    prs = Presentation(input_pptx)
    
    NAVY = RGBColor(0, 51, 102)        # #003366
    ACCENT_BLUE = RGBColor(22, 101, 192) # #1665C0
    DARK_TEXT = RGBColor(30, 41, 59)   # Slate 800
    MUTED_TEXT = RGBColor(71, 85, 105) # Slate 600
    CARD_BG = RGBColor(248, 250, 252)

    def set_para(p, text, font_size=12, bold=False, color=DARK_TEXT, space_after=4):
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(space_after)

    def style_title(shape, title_text):
        if shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.name = "Arial"
            p.font.size = Pt(22)
            p.font.bold = True
            p.font.color.rgb = NAVY

    def embed_diagram(slide, img_filename, box_left=Inches(0.8), box_top=Inches(1.5), box_w=Inches(8.4), box_h=Inches(5.4)):
        img_path = os.path.join(updir, img_filename)
        if not os.path.exists(img_path):
            print(f"Warning: Image {img_filename} not found!")
            return
        left, top, w, h = fit_image_box(img_path, box_left, box_top, box_w, box_h)
        
        # Subtle card border behind the diagram for clean presentation
        bg_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left - Inches(0.05), top - Inches(0.05), w + Inches(0.1), h + Inches(0.1))
        bg_card.fill.solid()
        bg_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        bg_card.line.color.rgb = RGBColor(210, 220, 235)
        bg_card.line.width = Pt(1.5)

        # Add image
        slide.shapes.add_picture(img_path, left, top, w, h)

    # ----------------------------------------------------
    # SLIDE 1: Cover
    # ----------------------------------------------------
    s1 = prs.slides[0]
    cover_box = s1.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(8.4), Inches(2.8))
    tf1 = cover_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    set_para(p, "Smart University Complaint System", font_size=26, bold=True, color=NAVY, space_after=6)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    set_para(p2, "A Web-Based Platform for Digital Complaint Management", font_size=14, bold=True, color=ACCENT_BLUE, space_after=8)
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf1.add_paragraph()
    set_para(p3, "Group Id: S26PROJECT4EE01 (BC220400155, BC220400156)", font_size=12, bold=True, color=DARK_TEXT, space_after=4)
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf1.add_paragraph()
    set_para(p4, "Supervisor Name: Muhammad Raheel  |  Virtual University of Pakistan", font_size=11, bold=False, color=MUTED_TEXT, space_after=2)
    p4.alignment = PP_ALIGN.CENTER


    # ----------------------------------------------------
    # SLIDE 2: Project Title (with VU Logo)
    # ----------------------------------------------------
    s2 = prs.slides[1]
    title_shape_s2 = s2.shapes[0]
    tf2 = title_shape_s2.text_frame
    tf2.word_wrap = True
    tf2.clear()

    p = tf2.paragraphs[0]
    set_para(p, "Smart University Complaint System (SUCS)", font_size=25, bold=True, color=NAVY, space_after=8)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf2.add_paragraph()
    set_para(p2, "Final Project Presentation | CS619", font_size=14, bold=True, color=ACCENT_BLUE, space_after=10)
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf2.add_paragraph()
    set_para(p3, "Group Id: S26PROJECT4EE01\nGroup Members: BC220400155, BC220400156", font_size=12, bold=True, color=DARK_TEXT, space_after=6)
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf2.add_paragraph()
    set_para(p4, "Project Supervisor: Muhammad Raheel", font_size=11, bold=False, color=MUTED_TEXT, space_after=2)
    p4.alignment = PP_ALIGN.CENTER


    # ----------------------------------------------------
    # SLIDE 3: Group Members Introduction
    # ----------------------------------------------------
    s3 = prs.slides[2]
    style_title(s3.shapes[0], "Group Members Introduction")
    body_s3 = s3.shapes[1]
    body_s3.top = Inches(1.8)
    body_s3.height = Inches(5.2)
    tf3 = body_s3.text_frame
    tf3.word_wrap = True
    tf3.clear()

    members = [
        ("Group Member 1: BC220400155", 
         "Lead Software Architecture, Python Flask Backend Engine, SQLite 3 Relational Database Normalization, Automated Keyword Priority Triage, and Conflict-of-Interest Firewall Enforcement."),
        ("Group Member 2: BC220400156", 
         "UI/UX Engineering, Jinja2 Template Integration, Tailwind CSS Responsive Interface Design, Student Dashboard & Quota Counter, and Pytest Automated Quality Assurance Test Suite."),
        ("Project Supervisor: Muhammad Raheel", 
         "Academic supervision, requirements validation, methodology oversight, and evaluation guidance for the Virtual University Final Project (CS619).")
    ]

    for i, (m_title, m_desc) in enumerate(members):
        p_hdr = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        set_para(p_hdr, f"• {m_title}", font_size=12, bold=True, color=NAVY, space_after=2)
        p_d = tf3.add_paragraph()
        set_para(p_d, f"    Role & Scope: {m_desc}", font_size=10, bold=False, color=DARK_TEXT, space_after=10)


    # ----------------------------------------------------
    # SLIDE 4: INTRODUCTION OF PROJECT
    # ----------------------------------------------------
    s4 = prs.slides[3]
    style_title(s4.shapes[0], "INTRODUCTION OF PROJECT")

    if len(s4.shapes) > 1 and s4.shapes[1].has_text_frame and s4.shapes[1].text_frame.text.strip() == "":
        s4.shapes[1].text_frame.clear()

    tb_s4 = s4.shapes[3] if len(s4.shapes) > 3 and s4.shapes[3].has_text_frame else s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.2))
    tf4 = tb_s4.text_frame
    tf4.word_wrap = True
    tf4.clear()

    intro_points = [
        ("Problem Statement", "Traditional campus grievance methods (paper suggestion boxes, unstructured emails) result in lost complaints, zero traceability, spamming, and severe conflict of interest when staff handle complaints filed against themselves."),
        ("Proposed Solution", "SUCS is an intelligent, role-governed web application featuring dedicated portals for Students, Administrators, and the Head of Department (HOD) with automated routing and status monitoring."),
        ("Intelligent Triage & Automation", "Includes automated hazard keyword priority escalation (High P1), 5-complaint monthly quotas, Conflict-of-Interest isolation (direct HOD escalation), and Category-Target problem clustering for 1-Click batch resolution."),
        ("Project Scope & Boundaries", "Focuses on internal web grievance management. Mobile apps, external SMS/Email gateways, and third-party ERP integration are scoped as future iterations.")
    ]

    for i, (p_t, p_d) in enumerate(intro_points):
        p_hdr = tf4.paragraphs[0] if i == 0 else tf4.add_paragraph()
        set_para(p_hdr, f"▶ {p_t}:", font_size=11.5, bold=True, color=NAVY, space_after=1)
        p_desc = tf4.add_paragraph()
        set_para(p_desc, f"    • {p_d}", font_size=10, bold=False, color=DARK_TEXT, space_after=6)


    # ----------------------------------------------------
    # SLIDE 5: Process Model (EMBED DIAGRAM)
    # ----------------------------------------------------
    s5 = prs.slides[4]
    # Replace placeholder text box with title shape and image
    tb_s5 = s5.shapes[0]
    tb_s5.left = Inches(0.6)
    tb_s5.top = Inches(0.3)
    tb_s5.width = Inches(8.8)
    tb_s5.height = Inches(1.0)
    tf5 = tb_s5.text_frame
    tf5.word_wrap = True
    tf5.clear()
    p_t = tf5.paragraphs[0]
    set_para(p_t, "Process Model: VU Hybrid Waterfall-Spiral Lifecycle", font_size=20, bold=True, color=NAVY, space_after=2)
    p_sub = tf5.add_paragraph()
    set_para(p_sub, "Waterfall phase discipline integrated with Spiral risk mitigation reviews", font_size=10.5, bold=True, color=ACCENT_BLUE, space_after=2)

    # Embed Process Model Diagram from SRS: media_1786788731168.png
    embed_diagram(s5, "media_1786788731168.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 6: Use Case Diagram (EMBED DIAGRAM)
    # ----------------------------------------------------
    s6 = prs.slides[5]
    style_title(s6.shapes[0], "Use Case Diagram")
    # Clear text placeholder
    s6.shapes[1].text_frame.clear()

    # Embed Use Case Diagram from SRS: media_1786788688397.png
    embed_diagram(s6, "media_1786788688397.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 7: Business Rules Catalog
    # ----------------------------------------------------
    s7 = prs.slides[6]
    style_title(s7.shapes[0], "Business Rules Catalog (BRC)")
    body_s7 = s7.shapes[1]
    body_s7.top = Inches(1.8)
    body_s7.height = Inches(5.2)
    tf7 = body_s7.text_frame
    tf7.word_wrap = True
    tf7.clear()

    br_rules = [
        ("BR-01: Monthly Quota Limit (FR3)", "Each student is limited to exactly 5 complaints per calendar month to prevent queue flooding and frivolous submissions."),
        ("BR-02: Conflict of Interest (COI) Firewall (FR7)", "If a complaint is directed against Administration, it completely bypasses the Admin queue and routes directly to the HOD."),
        ("BR-03: Emergency Keyword Priority Escalation (FR6)", "Complaints indicating physical danger (urgent, fire, shock, leak, hazard, broken wire) automatically escalate to High Priority (P1)."),
        ("BR-04: Intelligent Auto-Grouping (FR9)", "Complaints regarding the same Category and Target are automatically clustered under a common group_id token."),
        ("BR-05: 1-Click Batch Resolution", "Admin and HOD can resolve all linked complaints in a cluster simultaneously with a single click and uniform audit note."),
        ("BR-06: Mandatory Progress Remarks", "Changing ticket status to 'In-Progress' strictly mandates an investigative comment to keep students informed.")
    ]

    for i, (r_t, r_d) in enumerate(br_rules):
        p_r = tf7.paragraphs[0] if i == 0 else tf7.add_paragraph()
        set_para(p_r, f"✔ {r_t}", font_size=11, bold=True, color=NAVY, space_after=1)
        p_desc = tf7.add_paragraph()
        set_para(p_desc, f"    • {r_d}", font_size=9.5, bold=False, color=DARK_TEXT, space_after=5)


    # ----------------------------------------------------
    # SLIDE 8: Architectural Design (EMBED DIAGRAM)
    # ----------------------------------------------------
    s8 = prs.slides[7]
    style_title(s8.shapes[0], "Architectural Design (3-Tier Layered MVC)")
    # Clear text placeholder
    if len(s8.shapes) > 1 and s8.shapes[1].has_text_frame:
        s8.shapes[1].text_frame.clear()

    # Embed Architecture Diagram from SRS: media_1786788922762.png
    embed_diagram(s8, "media_1786788922762.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 9: Team Structure
    # ----------------------------------------------------
    s9 = prs.slides[8]
    style_title(s9.shapes[0], "Team Structure & Work Breakdown")
    body_s9 = s9.shapes[1]
    body_s9.top = Inches(1.8)
    body_s9.height = Inches(5.2)
    tf9 = body_s9.text_frame
    tf9.word_wrap = True
    tf9.clear()

    team_alloc = [
        ("Group Member 1: BC220400155 (Backend & Architecture Lead)", [
            "Project architecture definition & 3-Tier Layered MVC specification.",
            "Python 3.9 Flask backend routing, session management, and Werkzeug PBKDF2 hashing.",
            "Normalized SQLite 3 database schema design (3NF), foreign keys, and indices.",
            "Emergency keyword priority triage engine and Conflict-of-Interest routing logic."
        ]),
        ("Group Member 2: BC220400156 (Frontend & Quality Assurance Lead)", [
            "Tailwind CSS aesthetic styling, responsive grid layouts, and color-coded status badges.",
            "Jinja2 template architecture (Student submission, Admin desk, HOD console, DB viewer).",
            "Software Requirements Specification (SRS) & Use Case Scenario documentation.",
            "Automated Pytest test suite implementation (9 automated test cases, 100% pass rate)."
        ])
    ]

    for i, (m_head, tasks) in enumerate(team_alloc):
        p_h = tf9.paragraphs[0] if i == 0 else tf9.add_paragraph()
        set_para(p_h, f"• {m_head}", font_size=11.5, bold=True, color=NAVY, space_after=2)
        for t in tasks:
            p = tf9.add_paragraph()
            set_para(p, f"    - {t}", font_size=9.5, bold=False, color=DARK_TEXT, space_after=3)


    # ----------------------------------------------------
    # SLIDE 10: Project Schedule (EMBED DIAGRAM)
    # ----------------------------------------------------
    s10 = prs.slides[9]
    style_title(s10.shapes[0], "Project Schedule (Gantt Chart)")
    if len(s10.shapes) > 1 and s10.shapes[1].has_text_frame:
        s10.shapes[1].text_frame.clear()

    # Embed Gantt Chart from SRS: media_1786788747833.png
    embed_diagram(s10, "media_1786788747833.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 11: Sequence Diagrams (EMBED DIAGRAM)
    # ----------------------------------------------------
    s11 = prs.slides[10]
    style_title(s11.shapes[0], "Sequence Diagram: User Login & Complaint Workflow")
    if len(s11.shapes) > 1 and s11.shapes[1].has_text_frame:
        s11.shapes[1].text_frame.clear()

    # Embed Sequence Diagram 1 from SRS: media_1786788895708.png
    embed_diagram(s11, "media_1786788895708.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 12: Logical Model / Class Diagram (EMBED DIAGRAM)
    # ----------------------------------------------------
    s12 = prs.slides[11]
    style_title(s12.shapes[0], "Logical Model / Class Diagram")

    # Embed Class Diagram from SRS: media_1786788954827.png
    embed_diagram(s12, "media_1786788954827.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 13: Entity-Relationship Diagram (EMBED DIAGRAM)
    # ----------------------------------------------------
    s13 = prs.slides[12]
    style_title(s13.shapes[0], "Entity-Relationship Diagram (ERD)")

    # Embed ERD Diagram from SRS: media_1786788850589.png
    embed_diagram(s13, "media_1786788850589.png", box_left=Inches(0.6), box_top=Inches(1.6), box_w=Inches(8.8), box_h=Inches(3.5))

    # Add explanatory description below ERD
    erd_desc_box = s13.shapes.add_textbox(Inches(0.6), Inches(5.3), Inches(8.8), Inches(1.8))
    tf_erd = erd_desc_box.text_frame
    tf_erd.word_wrap = True
    p1 = tf_erd.paragraphs[0]
    set_para(p1, "ERD Cardinality & Relational Design:", font_size=11, bold=True, color=NAVY, space_after=2)
    p2 = tf_erd.add_paragraph()
    set_para(p2, "• USER (1 : 0..1) STUDENT: 1-to-1 specialization for student users tracking monthly complaint quota count.", font_size=9.5, bold=False, color=DARK_TEXT, space_after=2)
    p3 = tf_erd.add_paragraph()
    set_para(p3, "• STUDENT (1 : N) COMPLAINT: Each student can submit multiple complaints (max 5 per calendar month).", font_size=9.5, bold=False, color=DARK_TEXT, space_after=2)
    p4 = tf_erd.add_paragraph()
    set_para(p4, "• DEPARTMENT (1 : N) COMPLAINT: Complaints are categorized by target department for isolated triage.", font_size=9.5, bold=False, color=DARK_TEXT, space_after=2)


    # ----------------------------------------------------
    # SLIDE 14: Database Diagram (EMBED DIAGRAM)
    # ----------------------------------------------------
    s14 = prs.slides[13]
    style_title(s14.shapes[0], "Database Diagram & Schema Specification")

    # Embed Database Schema Diagram from SRS: media_1786788977208.png
    embed_diagram(s14, "media_1786788977208.png", box_left=Inches(0.6), box_top=Inches(1.4), box_w=Inches(8.8), box_h=Inches(5.7))


    # ----------------------------------------------------
    # SLIDE 15: User Interfaces (EMBED GUI SCREENSHOTS)
    # ----------------------------------------------------
    s15 = prs.slides[14]
    style_title(s15.shapes[0], "User Interfaces (Application Mockups)")
    if len(s15.shapes) > 1 and s15.shapes[1].has_text_frame:
        s15.shapes[1].text_frame.clear()

    # Embed two prominent GUI screenshots side by side:
    # Left: Student Dashboard (media_1786788997712.png)
    # Right: Admin Queue (media_1786789010431.png)
    embed_diagram(s15, "media_1786788997712.png", box_left=Inches(0.6), box_top=Inches(1.5), box_w=Inches(4.2), box_h=Inches(5.4))
    embed_diagram(s15, "media_1786789010431.png", box_left=Inches(5.2), box_top=Inches(1.5), box_w=Inches(4.2), box_h=Inches(5.4))


    # ----------------------------------------------------
    # SLIDE 16: Tools
    # ----------------------------------------------------
    s16 = prs.slides[15]
    style_title(s16.shapes[0], "Tools & Technologies Stack")
    body_s16 = s16.shapes[1]
    body_s16.top = Inches(1.8)
    body_s16.height = Inches(5.2)
    tf16 = body_s16.text_frame
    tf16.word_wrap = True
    tf16.clear()

    tools_list = [
        ("Python 3.9+ Programming Language", "Chosen for rapid prototyping, robust standard library, security primitives, and clean syntax."),
        ("Flask 2.x Microframework", "Lightweight WSGI web framework providing granular control over HTTP routing, sessions, and request dispatching."),
        ("SQLite 3 Relational Database Engine", "Zero-configuration, serverless, self-contained relational database guaranteeing ACID transaction compliance."),
        ("Werkzeug Security Module", "Industry-grade password cryptography implementing salted PBKDF2-SHA256 password hashing."),
        ("Tailwind CSS Frontend Styling", "Modern utility-first CSS framework delivering clean, accessible responsive cards, modals, and status badges."),
        ("Jinja2 Template Engine", "Server-side template engine providing template inheritance, safe variable escaping, and role-based view isolation."),
        ("Pytest 8.x Testing Framework", "Modern automated testing framework driving 9 automated integration test cases verifying 100% of business rules.")
    ]

    for i, (tool_n, tool_d) in enumerate(tools_list):
        p_t = tf16.paragraphs[0] if i == 0 else tf16.add_paragraph()
        set_para(p_t, f"✔ {tool_n}", font_size=11, bold=True, color=NAVY, space_after=1)
        p_desc = tf16.add_paragraph()
        set_para(p_desc, f"    • {tool_d}", font_size=9.5, bold=False, color=DARK_TEXT, space_after=4)

    prs.save(output_pptx)
    print(f"Presentation with embedded diagrams saved to: {output_pptx}")

if __name__ == "__main__":
    up_dir = "/Users/usman/.gemini/antigravity/brain/05b706b9-e3c5-4023-b5b0-9f886b725595/.user_uploaded"
    in_pptx = "/Users/usman/Desktop/UniversityComplaintManagement/original_template.pptx"
    out_pptx = "/Users/usman/Desktop/UniversityComplaintManagement/SUCS_VU_Template_Presentation.pptx"
    build_presentation_with_diagrams(in_pptx, out_pptx, up_dir)
